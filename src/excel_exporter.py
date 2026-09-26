from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def export_invoice_to_excel(
    invoice: dict[str, Any],
    output_path: Path,
) -> None:
    """Export validated invoice data to an Excel workbook."""

    workbook = Workbook()

    # -------------------------------------------------
    # Sheet 1: Invoice Summary
    # -------------------------------------------------
    summary_sheet = workbook.active
    summary_sheet.title = "Invoice Summary"

    summary_data = [
        ("Invoice Number", invoice.get("invoice_number")),
        ("Invoice Date", invoice.get("invoice_date")),
        ("Due Date", invoice.get("due_date")),
        ("Payment Terms", invoice.get("payment_terms")),
        ("PO Number", invoice.get("po_number")),
        ("Currency", invoice.get("currency")),
        ("Supplier Name", invoice.get("supplier_name")),
        ("Supplier GSTIN", invoice.get("supplier_gstin")),
        ("Customer Name", invoice.get("customer_name")),
        ("Customer GSTIN", invoice.get("customer_gstin")),
        ("Subtotal", invoice.get("subtotal")),
        ("GST", invoice.get("gst")),
        ("Total", invoice.get("total")),
    ]

    summary_sheet.append(["Field", "Value"])

    for cell in summary_sheet[1]:
        cell.font = Font(bold=True)

    for row in summary_data:
        summary_sheet.append(row)

    # -------------------------------------------------
    # Sheet 2: Line Items
    # -------------------------------------------------
    items_sheet = workbook.create_sheet("Line Items")

    headers = [
        "Item Number",
        "Description",
        "Quantity",
        "Unit Price",
        "Amount",
    ]

    items_sheet.append(headers)

    for cell in items_sheet[1]:
        cell.font = Font(bold=True)

    for item in invoice.get("items", []):
        items_sheet.append(
            [
                item.get("item_number"),
                item.get("description"),
                item.get("quantity"),
                item.get("unit_price"),
                item.get("amount"),
            ]
        )

    # -------------------------------------------------
    # Auto-size columns
    # -------------------------------------------------
    for sheet in workbook.worksheets:
        for column_cells in sheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

            sheet.column_dimensions[column_letter].width = max_length + 2

    # -------------------------------------------------
    # Save workbook
    # -------------------------------------------------
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)